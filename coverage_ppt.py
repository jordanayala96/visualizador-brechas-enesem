from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Iterable

import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from coverage_data import ReportData


PRIMARY = "2C70E7"
NAVY = "17365D"
CYAN = "00A6C8"
AMBER = "D99A2B"
CORAL = "C75252"
SLATE = "5F6B7A"
PALE_BLUE = "EAF2FF"
PALE_CYAN = "E8F7FA"
LIGHT = "F5F7FA"
GRID = "D7E0EA"
WHITE = "FFFFFF"
BLACK = "1F2937"

FONT = "Century Gothic"
SLIDE_W = 13.333
SLIDE_H = 7.5


MONTHS_ES = {
    1: "enero", 2: "febrero", 3: "marzo", 4: "abril", 5: "mayo", 6: "junio",
    7: "julio", 8: "agosto", 9: "septiembre", 10: "octubre",
    11: "noviembre", 12: "diciembre",
}


def _rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value.replace("#", ""))


def _asset(assets_dir: str | Path, name: str) -> Path:
    path = Path(assets_dir) / name
    if not path.exists():
        raise FileNotFoundError(f"No se encontró el recurso gráfico {path}.")
    return path


def _background(slide, image_path: Path) -> None:
    slide.shapes.add_picture(str(image_path), 0, 0, width=Inches(SLIDE_W), height=Inches(SLIDE_H))


def _add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 18,
    color: str = BLACK,
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    valign: MSO_ANCHOR = MSO_ANCHOR.MIDDLE,
    margin: float = 0.04,
) -> object:
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.text = str(text)
    paragraph.alignment = align
    paragraph.font.name = FONT
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = _rgb(color)
    return shape


def _add_content_title(slide, title: str) -> None:
    _add_text(slide, title, 1.18, 0.30, 10.35, 0.46, size=24, color=BLACK, bold=True)
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(1.18), Inches(0.84), Inches(1.1), Inches(0.055)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = _rgb(PRIMARY)
    line.line.fill.background()


def _format_number(value: object, percent: bool = False) -> str:
    if pd.isna(value):
        return "—"
    if percent:
        return f"{float(value):.1%}"
    if isinstance(value, (float, np.floating)) and not float(value).is_integer():
        return f"{float(value):,.1f}"
    return f"{int(round(float(value))):,}"


def _cell_fill(cell, color: str) -> None:
    cell.fill.solid()
    cell.fill.fore_color.rgb = _rgb(color)


def _cell_text(
    cell,
    text: object,
    *,
    size: float,
    color: str,
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.CENTER,
) -> None:
    cell.text = str(text)
    cell.margin_left = Inches(0.025)
    cell.margin_right = Inches(0.025)
    cell.margin_top = Inches(0.015)
    cell.margin_bottom = Inches(0.015)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    for paragraph in cell.text_frame.paragraphs:
        paragraph.alignment = align
        paragraph.font.name = FONT
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = _rgb(color)


def _add_table(
    slide,
    frame: pd.DataFrame,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    percent_columns: Iterable[str] = (),
    header_size: float = 8.0,
    body_size: float = 8.0,
    first_column_width: float | None = None,
    title: str | None = None,
) -> object:
    if title:
        _add_text(slide, title, x, y - 0.30, w, 0.26, size=11, color=NAVY, bold=True)
    rows, cols = len(frame) + 1, len(frame.columns)
    shape = slide.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(h))
    table = shape.table
    if first_column_width is not None and cols > 1:
        table.columns[0].width = Inches(first_column_width)
        remaining = (w - first_column_width) / (cols - 1)
        for index in range(1, cols):
            table.columns[index].width = Inches(remaining)
    for index, column in enumerate(frame.columns):
        cell = table.cell(0, index)
        _cell_fill(cell, NAVY)
        _cell_text(cell, column, size=header_size, color=WHITE, bold=True)
    percent_set = set(percent_columns)
    for row_index, (_, row) in enumerate(frame.iterrows(), start=1):
        is_total = str(row.iloc[0]).strip().upper() in {
            "NACIONAL", "TOTAL", "TOTAL ENVIADAS", "TOTAL RECIBIDAS"
        }
        for col_index, column in enumerate(frame.columns):
            cell = table.cell(row_index, col_index)
            if is_total:
                fill, color, bold = PRIMARY, WHITE, True
            elif row_index % 2:
                fill, color, bold = WHITE, BLACK, col_index == 0
            else:
                fill, color, bold = LIGHT, BLACK, col_index == 0
            _cell_fill(cell, fill)
            value = row[column]
            if column in percent_set:
                text = _format_number(value, percent=True)
            elif isinstance(value, (int, float, np.number)) and not isinstance(value, bool):
                text = _format_number(value)
            else:
                text = "" if pd.isna(value) else str(value)
            _cell_text(
                cell,
                text,
                size=body_size,
                color=color,
                bold=bold,
                align=PP_ALIGN.LEFT if col_index == 0 else PP_ALIGN.CENTER,
            )
    return table


def _add_coverage_chart(
    slide,
    frame: pd.DataFrame,
    x: float,
    y: float,
    w: float,
    h: float,
    executed_column: str,
    title: str,
) -> None:
    chart_data = ChartData()
    chart_data.categories = [f"S{int(week)}" for week in frame["Semana"]]
    chart_data.add_series(
        "Muestra pendiente",
        [None if pd.isna(value) else float(value) for value in frame["Muestra pendiente según planificación"]],
    )
    chart_data.add_series(
        "Planificado acumulado",
        [None if pd.isna(value) else float(value) for value in frame["Planificado acumulado"]],
    )
    chart_data.add_series(
        "Ejecutado acumulado",
        [None if pd.isna(value) else float(value) for value in frame[executed_column]],
    )
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.LINE_MARKERS,
        Inches(x), Inches(y), Inches(w), Inches(h),
        chart_data,
    ).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = title
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.TOP
    chart.legend.include_in_layout = False
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.minimum_scale = 0
    chart.category_axis.tick_labels.font.size = Pt(7)
    chart.category_axis.tick_labels.font.name = FONT
    chart.value_axis.tick_labels.font.size = Pt(8)
    chart.value_axis.tick_labels.font.name = FONT
    chart.chart_title.text_frame.paragraphs[0].font.name = FONT
    chart.chart_title.text_frame.paragraphs[0].font.size = Pt(11)
    chart.chart_title.text_frame.paragraphs[0].font.bold = True
    chart.chart_title.text_frame.paragraphs[0].font.color.rgb = _rgb(NAVY)
    colors = [SLATE, CYAN, PRIMARY]
    widths = [1.4, 2.0, 2.6]
    for series, color, width in zip(chart.series, colors, widths):
        series.format.line.color.rgb = _rgb(color)
        series.format.line.width = Pt(width)


def _coverage_table_for_slide(frame: pd.DataFrame, executed: str, effective: str) -> pd.DataFrame:
    result = frame[[
        "Coordinación Zonal", "Muestra", "Muestra actual", "Planificadas",
        "% planificadas", executed, f"% {executed.lower()}", effective,
        f"% {effective.lower()}",
    ]].copy()
    return result.rename(columns={
        "% planificadas": "% planif. nac.",
        f"% {executed.lower()}": "% ejec. nac.",
        f"% {effective.lower()}": "% efect. nac.",
    })


def _novelty_display(frame: pd.DataFrame, zones: list[str]) -> pd.DataFrame:
    columns = ["Novedad", *zones, "NACIONAL"]
    if frame.empty:
        return pd.DataFrame([["Sin novedades para los criterios seleccionados", *([0] * (len(columns) - 1))]], columns=columns)
    rows: list[dict[str, object]] = []
    for section, part in frame.groupby("Sección", sort=False):
        for _, row in part.iterrows():
            item = {column: row.get(column, 0) for column in columns}
            item["Novedad"] = row["Novedad"]
            rows.append(item)
        subtotal: dict[str, object] = {"Novedad": f"Total {section.lower()}"}
        for column in [*zones, "NACIONAL"]:
            subtotal[column] = int(pd.to_numeric(part[column], errors="coerce").fillna(0).sum())
        rows.append(subtotal)
    total: dict[str, object] = {"Novedad": "TOTAL"}
    for column in [*zones, "NACIONAL"]:
        total[column] = int(pd.to_numeric(frame[column], errors="coerce").fillna(0).sum())
    rows.append(total)
    return pd.DataFrame(rows, columns=columns)


def _add_novelty_table(slide, frame: pd.DataFrame, x: float, y: float, w: float, h: float) -> None:
    rows, cols = len(frame) + 1, len(frame.columns)
    shape = slide.shapes.add_table(rows, cols, Inches(x), Inches(y), Inches(w), Inches(h))
    table = shape.table
    first_width = 7.0
    table.columns[0].width = Inches(first_width)
    for index in range(1, cols):
        table.columns[index].width = Inches((w - first_width) / (cols - 1))
    body_size = 7.5 if len(frame) <= 16 else 6.7
    for index, column in enumerate(frame.columns):
        _cell_fill(table.cell(0, index), NAVY)
        _cell_text(table.cell(0, index), column, size=8.0, color=WHITE, bold=True)
    for row_index, (_, row) in enumerate(frame.iterrows(), start=1):
        label = str(row.iloc[0])
        is_grand_total = label == "TOTAL"
        is_subtotal = label.lower().startswith("total ") and not is_grand_total
        for col_index, column in enumerate(frame.columns):
            cell = table.cell(row_index, col_index)
            if is_grand_total:
                fill, color, bold = NAVY, WHITE, True
            elif is_subtotal:
                fill, color, bold = PRIMARY, WHITE, True
            else:
                fill, color, bold = (WHITE if row_index % 2 else LIGHT), BLACK, col_index == 0
            _cell_fill(cell, fill)
            value = row[column]
            text = _format_number(value) if col_index else str(value)
            _cell_text(
                cell,
                text,
                size=body_size,
                color=color,
                bold=bold,
                align=PP_ALIGN.LEFT if col_index == 0 else PP_ALIGN.CENTER,
            )


def generate_coverage_pptx(
    report: ReportData,
    assets_dir: str | Path,
    report_month: int | None = None,
    report_year: int | None = None,
) -> bytes:
    assets_dir = Path(assets_dir)
    report_year = int(report_year or report.cutoff.year)
    report_month = int(report_month or report.cutoff.month)
    enesem_year = report_year - 1

    presentation = Presentation()
    presentation.slide_width = Inches(SLIDE_W)
    presentation.slide_height = Inches(SLIDE_H)
    blank = presentation.slide_layouts[6]

    cover = presentation.slides.add_slide(blank)
    _background(cover, _asset(assets_dir, "cover_enesem.png"))
    _add_text(cover, "Encuesta Estructural\nEmpresarial", 0.78, 1.72, 5.75, 1.15, size=30, color=SLATE, bold=True)
    _add_text(cover, f"ENESEM {enesem_year}", 1.55, 3.25, 3.9, 0.5, size=21, color=PRIMARY, bold=True)
    _add_text(
        cover,
        f"{MONTHS_ES[report_month].capitalize()} - {report_year}",
        0.95, 4.20, 3.65, 0.42,
        size=15, color=NAVY, bold=True, align=PP_ALIGN.CENTER,
    )

    separator = presentation.slides.add_slide(blank)
    _background(separator, _asset(assets_dir, "separator_enesem.png"))
    date_text = f"{report.cutoff.day} de {MONTHS_ES[report.cutoff.month]} de {report.cutoff.year}"
    _add_text(
        separator,
        f"Reporte de cobertura\ncon corte al {date_text}",
        3.05, 2.65, 7.45, 1.55,
        size=27, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
    )

    slide = presentation.slides.add_slide(blank)
    _background(slide, _asset(assets_dir, "content_enesem.png"))
    _add_content_title(slide, "Cobertura de socialización y diligenciamiento")
    left = report.socialization[[
        "Coordinación Zonal", "Muestra original", "Enviadas", "Recibidas", "Muestra actual", "% muestra",
    ]]
    right = report.socialization[[
        "Coordinación Zonal", "Socializadas", "% socialización", "Diligenciadas", "% diligenciamiento", "Por socializar",
    ]]
    _add_table(
        slide, left, 0.52, 1.18, 6.15, 2.10,
        percent_columns=["% muestra"], first_column_width=1.35,
        title="Distribución de la muestra",
    )
    _add_table(
        slide, right, 6.82, 1.18, 6.0, 2.10,
        percent_columns=["% socialización", "% diligenciamiento"], first_column_width=1.35,
        title="Avance de socialización y diligenciamiento",
    )
    _add_table(
        slide, report.transfers, 0.52, 3.72, 12.26, 2.55,
        first_column_width=2.25, title="Empresas que informan en otra Coordinación Zonal",
    )
    _add_text(
        slide,
        f"Semana seleccionada: {report.selected_week}",
        0.55, 6.98, 4.8, 0.22, size=8, color=SLATE,
    )

    for title, coverage, indicators, chart, executed, effective in (
        (
            "Cobertura de levantamiento", report.field_coverage, report.field_indicators,
            report.field_chart, "Ejecutadas", "Efectivas",
        ),
        (
            "Cobertura de crítica", report.critique_coverage, report.critique_indicators,
            report.critique_chart, "Criticadas", "Efectivas criticadas",
        ),
    ):
        phase_slide = presentation.slides.add_slide(blank)
        _background(phase_slide, _asset(assets_dir, "content_enesem.png"))
        _add_content_title(phase_slide, title)
        coverage_display = _coverage_table_for_slide(coverage, executed, effective)
        _add_table(
            phase_slide, coverage_display, 0.52, 1.22, 7.75, 1.72,
            percent_columns=["% planif. nac.", "% ejec. nac.", "% efect. nac."],
            first_column_width=1.55, title="Cobertura acumulada por Coordinación Zonal",
            header_size=6.5, body_size=6.9,
        )
        _add_table(
            phase_slide, indicators, 8.45, 1.22, 4.35, 1.72,
            percent_columns=list(indicators.columns[1:]), first_column_width=1.25,
            title="Indicadores de cobertura", header_size=6.2, body_size=6.9,
        )
        executed_chart_column = executed + " acumuladas"
        chart_title = (
            "Cobertura acumulada de levantamiento nacional"
            if executed == "Ejecutadas"
            else "Cobertura acumulada de crítica nacional"
        )
        _add_coverage_chart(
            phase_slide, chart, 0.72, 3.20, 11.95, 3.45,
            executed_chart_column, chart_title,
        )
        _add_text(
            phase_slide,
            f"Planificación y resultados hasta la semana {report.selected_week}.",
            0.72, 6.76, 8.4, 0.25, size=8.5, color=SLATE,
        )

    progress_slide = presentation.slides.add_slide(blank)
    _background(progress_slide, _asset(assets_dir, "content_enesem.png"))
    _add_content_title(progress_slide, "Estado de avance de cobertura")
    status_display = report.progress_status.copy()
    if int(status_display["Sin registro de avance"].sum()) == 0:
        status_display = status_display.drop(columns="Sin registro de avance")
    _add_table(
        progress_slide, status_display, 0.75, 1.42, 5.75, 3.55,
        first_column_width=1.35, title="Estado de llenado",
    )
    _add_table(
        progress_slide, report.progress_bands, 6.85, 1.42, 5.70, 3.55,
        first_column_width=1.35, title="Avance de las empresas en proceso",
    )

    novelty_display = _novelty_display(report.novelties, report.zones)
    page_size = 17
    pages = max(1, int(np.ceil(len(novelty_display) / page_size)))
    for page in range(pages):
        novelty_slide = presentation.slides.add_slide(blank)
        _background(novelty_slide, _asset(assets_dir, "content_enesem.png"))
        page_title = "Novedades" if page == 0 else "Novedades (continuación)"
        _add_content_title(novelty_slide, page_title)
        part = novelty_display.iloc[page * page_size:(page + 1) * page_size].copy()
        _add_novelty_table(novelty_slide, part, 0.65, 1.22, 12.0, 5.65)
        _add_text(
            novelty_slide,
            "”. "
            "",
            0.70, 6.98, 11.5, 0.23, size=8.2, color=SLATE,
        )

    closing = presentation.slides.add_slide(blank)
    _background(closing, _asset(assets_dir, "closing_enesem.png"))
    _add_text(closing, "", 7.25, 2.55, 3.85, 0.75, size=34, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    _add_text(
        closing,
        "",
        6.72, 3.38, 4.90, 0.42, size=17, color=PRIMARY, bold=True, align=PP_ALIGN.CENTER,
    )

    output = BytesIO()
    presentation.save(output)
    return output.getvalue()
